from __future__ import annotations

import os
import shutil
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from docx.oxml import parse_xml
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from PIL import Image, ImageDraw, ImageFont

from rag_catalog.core.extractors.embedded_media import (
    count_embedded_images,
    extract_embedded_documents,
    iter_embedded_images,
    ocr_embedded_images,
)
from rag_catalog.core.extractors.files import (
    extract_docx,
    extract_docx_document,
    extract_pptx_document,
    extract_spreadsheet_document,
    extract_xlsx_document,
)
from rag_catalog.core.ocr_runtime import resolve_ocr_runtime

SAMPLE_TEXT = "ТЕСТ 123456"


def _resolve_tesseract() -> str:
    """Tesseract из RAG_TESSERACT_CMD / bundled tools/ (resolve_ocr_runtime) / PATH — без зашитых путей."""
    for candidate in (
        os.environ.get("RAG_TESSERACT_CMD", ""),
        resolve_ocr_runtime({}).get("tesseract_cmd", ""),
        shutil.which("tesseract") or "",
    ):
        if candidate and Path(candidate).exists():
            return candidate
    return ""


TESSERACT = _resolve_tesseract()
needs_tesseract = pytest.mark.skipif(not TESSERACT, reason="tesseract недоступен")


def _font(size: int):
    for name in ("arial.ttf", "arialbd.ttf", "DejaVuSans.ttf"):
        for root in (Path(r"C:\Windows\Fonts"), Path("/usr/share/fonts/truetype/dejavu")):
            candidate = root / name
            if candidate.exists():
                return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default(size=size)


def _text_png(text: str = SAMPLE_TEXT, size: tuple[int, int] = (1400, 400)) -> bytes:
    img = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(img)
    draw.text((40, 100), text, fill="black", font=_font(140))
    buf = BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _tiny_png() -> bytes:
    img = Image.new("RGB", (16, 16), "red")
    buf = BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _build_docx(tmp_path: Path, png: bytes, *, name: str = "sample.docx", body: str = "Абзац") -> Path:
    path = tmp_path / name
    doc = Document()
    if body:
        doc.add_paragraph(body)
    doc.add_picture(BytesIO(png))
    doc.save(path)
    return path


def _build_xlsx(tmp_path: Path, png: bytes, *, name: str = "sample.xlsx") -> Path:
    path = tmp_path / name
    wb = Workbook()
    ws = wb.active
    ws.title = "Данные"
    ws["A1"] = "ячейка"
    picture_path = tmp_path / f"{name}.png"
    picture_path.write_bytes(png)
    ws.add_image(XlImage(str(picture_path)), "B2")
    wb.save(path)
    return path


def _build_pptx(tmp_path: Path, png: bytes, *, name: str = "sample.pptx") -> Path:
    path = tmp_path / name
    slide_xml = """<?xml version="1.0" encoding="UTF-8"?>
<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
       xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
  <p:cSld><p:spTree><p:sp><p:txBody>
    <a:p><a:r><a:t>Слайд с картинкой</a:t></a:r></a:p>
  </p:txBody></p:sp></p:spTree></p:cSld>
</p:sld>
"""
    with ZipFile(path, "w", ZIP_DEFLATED) as zf:
        zf.writestr("ppt/slides/slide1.xml", slide_xml)
        zf.writestr("ppt/media/image1.png", png)
    return path


def test_iter_embedded_images_docx_yields_picture(tmp_path: Path) -> None:
    png = _text_png()
    path = _build_docx(tmp_path, png)

    diagnostics: dict = {}
    found = list(iter_embedded_images(path, diagnostics=diagnostics))

    assert len(found) == 1
    member, data = found[0]
    assert member.startswith("word/media/")
    assert data == png
    assert diagnostics["embedded_images"] == 1


def test_iter_embedded_images_xlsx_and_pptx(tmp_path: Path) -> None:
    png = _text_png()
    xlsx = _build_xlsx(tmp_path, png)
    pptx = _build_pptx(tmp_path, png)

    assert [m for m, _ in iter_embedded_images(xlsx)][0].startswith("xl/media/")
    assert [m for m, _ in iter_embedded_images(pptx)] == ["ppt/media/image1.png"]


def test_iter_embedded_images_skips_tiny_vector_and_limit(tmp_path: Path) -> None:
    path = tmp_path / "mixed.docx"
    big = _text_png()
    with ZipFile(path, "w", ZIP_DEFLATED) as zf:
        zf.writestr("word/media/image1.png", _tiny_png())  # < 8 KB
        zf.writestr("word/media/image2.emf", b"\x01" * 20000)  # vector
        zf.writestr("word/media/image3.png", big)
        zf.writestr("word/media/image4.png", big)
        # large bytes but small pixel size
        small_px = Image.new("RGB", (100, 100), "white")
        buf = BytesIO()
        small_px.save(buf, format="BMP")
        zf.writestr("word/media/image0.bmp", buf.getvalue() + b"\x00" * 9000)

    diagnostics: dict = {}
    found = list(iter_embedded_images(path, max_images=1, diagnostics=diagnostics))

    assert [m for m, _ in found] == ["word/media/image3.png"]
    assert diagnostics["embedded_images_skipped_small"] == 2
    assert diagnostics["embedded_images_skipped_vector"] == 1
    assert diagnostics["embedded_images_skipped_limit"] == 1
    assert count_embedded_images(path) == 2


def test_iter_embedded_images_ignores_non_ooxml_and_broken_zip(tmp_path: Path) -> None:
    txt = tmp_path / "plain.txt"
    txt.write_text("x", encoding="utf-8")
    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"not a zip")

    assert list(iter_embedded_images(txt)) == []
    assert list(iter_embedded_images(broken)) == []


def test_ocr_embedded_images_skip_ocr_reports_diagnostics(tmp_path: Path) -> None:
    path = _build_docx(tmp_path, _text_png())

    diagnostics: dict = {}
    blocks = ocr_embedded_images(path, skip_ocr=True, diagnostics=diagnostics)

    assert blocks == []
    assert diagnostics["embedded_images"] == 1
    assert diagnostics["ocr_skipped"] is True


def test_extract_docx_document_skip_ocr_metadata_marks_deferred(tmp_path: Path) -> None:
    path = _build_docx(tmp_path, _text_png(), body="")

    doc = extract_docx_document(path, skip_ocr=True)

    assert doc.blocks == ()
    assert doc.metadata["embedded_images"] == 1
    assert doc.metadata["ocr_skipped"] is True


def test_extract_docx_legacy_default_does_not_ocr(tmp_path: Path, monkeypatch) -> None:
    path = _build_docx(tmp_path, _text_png(), body="Только текст")
    monkeypatch.setattr(
        "rag_catalog.core.extractors.embedded_media._ocr_image_bytes",
        lambda *_a, **_k: (_ for _ in ()).throw(AssertionError("OCR must not run")),
    )

    assert extract_docx(path) == "Только текст"


def test_extract_docx_document_uses_fake_ocr(tmp_path: Path, monkeypatch) -> None:
    path = _build_docx(tmp_path, _text_png(), body="Абзац")
    monkeypatch.setattr(
        "rag_catalog.core.extractors.embedded_media._ocr_image_bytes",
        lambda *_a, **_k: "FAKE 999",
    )

    doc = extract_docx_document(path, skip_ocr=False)

    assert doc.blocks[0].text == "Абзац"
    image_blocks = [b for b in doc.blocks if b.section.startswith("image:")]
    assert len(image_blocks) == 1
    assert image_blocks[0].text == "FAKE 999"
    assert image_blocks[0].metadata["source"] == "embedded_image"
    assert doc.metadata["embedded_images"] == 1
    assert doc.metadata["embedded_images_ocr"] == 1


def test_extract_docx_document_reads_header_footer_textbox_footnotes(tmp_path: Path) -> None:
    path = tmp_path / "rich.docx"
    doc = Document()
    doc.add_paragraph("Тело документа")
    section = doc.sections[0]
    section.header.paragraphs[0].text = "Верхний колонтитул 111"
    section.footer.paragraphs[0].text = "Нижний колонтитул 222"
    textbox = parse_xml(
        '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        "<w:pict><w:txbxContent><w:p><w:r><w:t>Текстовое поле 333</w:t></w:r></w:p></w:txbxContent></w:pict></w:r>"
    )
    doc.paragraphs[0]._p.append(textbox)
    doc.save(path)
    footnotes = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:footnote w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
        '<w:footnote w:id="1"><w:p><w:r><w:t>Сноска 444</w:t></w:r></w:p></w:footnote>'
        "</w:footnotes>"
    )
    with ZipFile(path, "a", ZIP_DEFLATED) as zf:
        zf.writestr("word/footnotes.xml", footnotes)

    extracted = extract_docx_document(path, include_embedded=False)
    by_section = {b.section: b.text for b in extracted.blocks}

    assert "Тело документа" in by_section[""]
    assert by_section["header"] == "Верхний колонтитул 111"
    assert by_section["footer"] == "Нижний колонтитул 222"
    assert by_section["textbox"] == "Текстовое поле 333"
    assert "Сноска 444" in by_section["footnotes"]


def test_extract_embedded_documents_reads_nested_xlsx_and_skips_ole(tmp_path: Path) -> None:
    path = _build_docx(tmp_path, _text_png(), body="Внешний")
    wb = Workbook()
    wb.active["A1"] = "Вложенная таблица 555"
    inner = BytesIO()
    wb.save(inner)
    with ZipFile(path, "a", ZIP_DEFLATED) as zf:
        zf.writestr("word/embeddings/Лист Microsoft Excel.xlsx", inner.getvalue())
        zf.writestr("word/embeddings/oleObject1.bin", b"\xd0\xcf\x11\xe0" + b"\x00" * 100)

    diagnostics: dict = {}
    blocks = extract_embedded_documents(path, diagnostics=diagnostics)

    assert len(blocks) == 1
    assert blocks[0].text == "Вложенная таблица 555"
    assert blocks[0].section == "embedded:word/embeddings/Лист Microsoft Excel.xlsx"
    assert blocks[0].sheet == "Sheet"
    assert diagnostics["embedded_objects"] == 1
    assert diagnostics["embedded_objects_skipped"] == 1

    doc = extract_docx_document(path, skip_ocr=True)
    assert any("555" in b.text for b in doc.blocks)
    assert doc.metadata["embedded_objects"] == 1


def test_xlsx_and_pptx_documents_carry_embedded_diagnostics(tmp_path: Path) -> None:
    png = _text_png()
    xlsx = _build_xlsx(tmp_path, png)
    pptx = _build_pptx(tmp_path, png)

    xdoc = extract_spreadsheet_document(xlsx, skip_ocr=True)
    pdoc = extract_pptx_document(pptx, skip_ocr=True)

    assert xdoc.blocks[0].text == "ячейка"
    assert xdoc.metadata["embedded_images"] == 1 and xdoc.metadata["ocr_skipped"] is True
    assert pdoc.blocks[0].slide == 1
    assert pdoc.metadata["embedded_images"] == 1 and pdoc.metadata["ocr_skipped"] is True


@needs_tesseract
def test_ocr_embedded_images_recognizes_generated_text(tmp_path: Path) -> None:
    path = _build_docx(tmp_path, _text_png())

    diagnostics: dict = {}
    blocks = ocr_embedded_images(path, tesseract_cmd=TESSERACT, diagnostics=diagnostics)

    assert len(blocks) == 1
    assert "123456" in blocks[0].text
    assert blocks[0].section == "image:word/media/image1.png"
    assert diagnostics["embedded_images_ocr"] == 1


@needs_tesseract
def test_xlsx_document_ocr_of_embedded_picture(tmp_path: Path) -> None:
    path = _build_xlsx(tmp_path, _text_png())

    doc = extract_xlsx_document(path, tesseract_cmd=TESSERACT)

    assert any("123456" in b.text for b in doc.blocks if b.section.startswith("image:"))


@needs_tesseract
def test_docx_with_two_scanned_pictures_recognizes_both(tmp_path: Path) -> None:
    """Синтетический аналог «docx со сканами СТС»: два вставленных изображения с текстом."""
    path = tmp_path / "two-scans.docx"
    doc = Document()
    doc.add_picture(BytesIO(_text_png("СЕРИЯ 501049")))
    doc.add_picture(BytesIO(_text_png("НОМЕР 777123")))
    doc.save(path)

    extracted = extract_docx_document(path, tesseract_cmd=TESSERACT)

    image_text = "\n".join(b.text for b in extracted.blocks if b.section.startswith("image:"))
    assert extracted.metadata["embedded_images"] == 2
    assert "501049" in image_text.replace(" ", "")
    assert "777123" in image_text.replace(" ", "")
