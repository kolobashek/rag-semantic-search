"""Вложенные картинки/объекты OOXML в индексаторе (п. 1–2 ревью).

- docx идёт через extract_docx_document, OCR картинок управляется skip_ocr индексатора;
- xlsx/pptx получают skip_ocr / tesseract_cmd / include_embedded;
- при --no-ocr документ с картинками получает status=deferred_ocr (как скан PDF);
- явный block.section (image:…, header, footnotes) и block_kind попадают в payload.
"""

from __future__ import annotations

import random
import sys
import types
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from PIL import Image, ImageDraw

from index_rag import PAYLOAD_SCHEMA_VERSION, RAGIndexer
from rag_catalog.core import index_rag as index_rag_module
from rag_catalog.core.extractors import ExtractedDocument, TextBlock
from rag_catalog.core.index_state_db import IndexStateDB
from rag_catalog.core.indexing.ocr_deferral import (
    document_has_deferred_embedded_ocr,
    is_deferred_ocr_candidate,
)

# Длиннее min_chunk_chars (20): блок OCR короче отбрасывается чанкером как фрагмент.
SAMPLE_TEXT = "ТЕСТ 123456 ПАСПОРТ ТЕХНИКИ"


class _FakeVec:
    def __init__(self, vals):
        self.vals = vals

    def tolist(self):
        return list(self.vals)


class _FakeEmbedder:
    def encode(self, chunks, normalize_embeddings=True, batch_size=256, show_progress_bar=False):
        return [_FakeVec([0.1, 0.2, 0.3]) for _ in chunks]


class _FakeQdrant:
    def __init__(self) -> None:
        self.points_count = 0
        self.points: list = []

    def upsert(self, collection_name, points):
        self.points_count += len(points)
        self.points.extend(points)

    def get_collection(self, collection_name):
        return SimpleNamespace(points_count=self.points_count)


def _text_png(text: str = SAMPLE_TEXT, size: tuple[int, int] = (1200, 400)) -> bytes:
    """PNG с текстом (размер > 8 КБ и > 200 px, чтобы пройти фильтры iter_embedded_images)."""
    img = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(img)
    draw.text((40, 120), text, fill="black")
    # Полоса случайного шума: PNG её не сжимает, файл гарантированно > DEFAULT_MIN_IMAGE_BYTES (8 КБ).
    rng = random.Random(20260911)
    noise_h = 40
    noise = Image.frombytes("RGB", (size[0], noise_h), rng.randbytes(size[0] * noise_h * 3))
    img.paste(noise, (0, size[1] - noise_h))
    buf = BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _build_docx(path: Path, *, body: str = "Абзац документа для теста индексатора") -> Path:
    doc = Document()
    if body:
        doc.add_paragraph(body)
    doc.add_picture(BytesIO(_text_png()))
    doc.save(path)
    return path


def _make_indexer(tmp_path: Path, *, skip_ocr: bool = False) -> RAGIndexer:
    idx = RAGIndexer.__new__(RAGIndexer)
    idx.current_stage = "small"
    idx.catalog_path = tmp_path
    idx.collection_name = "catalog"
    idx.chunk_size = 500
    idx.chunk_overlap = 100
    idx.min_chunk_chars = 20
    idx.batch_size = 1000
    idx.max_chunks_per_file = 0
    idx.read_workers = 1
    idx.metadata_only_extensions = set()
    idx.state_db = IndexStateDB(str(tmp_path / "state" / "index_state.db"))
    idx.point_count = 0
    idx.payload_schema_version = PAYLOAD_SCHEMA_VERSION
    idx.run_id = ""
    idx._run_deleted_files = 0
    idx.small_office_mb = 20.0
    idx.small_pdf_mb = 2.0
    idx.skip_ocr = skip_ocr
    idx.index_embedded_media = True
    idx.ocr_tesseract_cmd = ""
    idx.synonym_map = {}
    idx.embedder = _FakeEmbedder()
    idx.qdrant = _FakeQdrant()
    idx._delete_file_vectors = lambda _p, **_k: None
    idx._cleanup_deleted_files = lambda _files, **_k: 0
    return idx


@pytest.fixture
def fake_ocr(monkeypatch):
    """pytesseract.image_to_string → SAMPLE_TEXT (без реального tesseract)."""
    calls: list = []
    try:
        import pytesseract  # type: ignore  # noqa: PLC0415
    except ImportError:
        pytesseract = types.ModuleType("pytesseract")
        pytesseract.pytesseract = types.SimpleNamespace(tesseract_cmd="")
        monkeypatch.setitem(sys.modules, "pytesseract", pytesseract)

    def _image_to_string(image, **kwargs):
        calls.append(kwargs)
        return SAMPLE_TEXT + "\n"

    monkeypatch.setattr(pytesseract, "image_to_string", _image_to_string, raising=False)
    return calls


def _content_payloads(idx: RAGIndexer) -> list[dict]:
    return [p.payload for p in idx.qdrant.points if str(p.payload.get("type") or "") != "file_metadata"]


# ── п.1: docx через extract_docx_document, OCR картинок по skip_ocr индексатора ──


def test_stage_runner_indexes_docx_picture_text_via_ocr(tmp_path: Path, fake_ocr) -> None:
    path = _build_docx(tmp_path / "scan.docx")
    idx = _make_indexer(tmp_path, skip_ocr=False)

    stats = idx.index_directory(stage="small")

    assert stats["error_files"] == 0
    assert len(fake_ocr) == 1
    payloads = _content_payloads(idx)
    image_payloads = [p for p in payloads if str(p.get("section") or "").startswith("image:")]
    assert len(image_payloads) == 1
    assert SAMPLE_TEXT in image_payloads[0]["text"]
    assert image_payloads[0]["section"] == "image:word/media/image1.png"
    assert image_payloads[0]["block_kind"] == "embedded_image"
    assert image_payloads[0]["provenance"]["section"] == "image:word/media/image1.png"
    body = [p for p in payloads if p.get("block_kind") == "body"]
    assert body and "Абзац документа" in body[0]["text"]
    row = idx.state_db.get_entry(str(path))
    assert row["stage"] == "content"
    assert row["status"] == "ok"


def test_stage_runner_no_ocr_docx_with_pictures_is_deferred_and_skipped_until_ocr(
    tmp_path: Path, fake_ocr
) -> None:
    path = _build_docx(tmp_path / "scan.docx")
    idx = _make_indexer(tmp_path, skip_ocr=True)

    first = idx.index_directory(stage="small")

    assert first["error_files"] == 0
    assert fake_ocr == []  # OCR не запускался
    row = idx.state_db.get_entry(str(path))
    assert row["status"] == "deferred_ocr"
    assert row["stage"] == "metadata"
    assert row["last_error"] == "deferred_ocr"
    # Текст тела документа при этом проиндексирован.
    assert any("Абзац документа" in p["text"] for p in _content_payloads(idx))

    second = idx.index_directory(stage="small")
    assert second["skipped_files"] == 1

    # Прогон с OCR дочитывает картинку.
    idx.skip_ocr = False
    third = idx.index_directory(stage="large")
    assert third["error_files"] == 0
    assert len(fake_ocr) == 1
    row = idx.state_db.get_entry(str(path))
    assert row["status"] == "ok"
    assert row["stage"] == "content"
    assert any(str(p.get("section") or "").startswith("image:") for p in _content_payloads(idx))


def test_docx_without_pictures_under_no_ocr_is_not_deferred(tmp_path: Path, fake_ocr) -> None:
    path = tmp_path / "plain.docx"
    doc = Document()
    doc.add_paragraph("Просто текст без картинок для индексатора")
    doc.save(path)
    idx = _make_indexer(tmp_path, skip_ocr=True)

    idx.index_directory(stage="small")

    row = idx.state_db.get_entry(str(path))
    assert row["status"] == "ok"
    assert row["stage"] == "content"


def test_process_file_marks_docx_with_pictures_deferred_under_no_ocr(tmp_path: Path, fake_ocr) -> None:
    path = _build_docx(tmp_path / "queued.docx")
    idx = _make_indexer(tmp_path, skip_ocr=True)
    idx.qdrant_timeout_sec = 5

    idx.process_file(path)

    row = idx.state_db.get_entry(str(path))
    assert row["status"] == "deferred_ocr"
    assert row["stage"] == "metadata"
    assert fake_ocr == []


def test_process_file_indexes_docx_picture_text_with_ocr(tmp_path: Path, fake_ocr) -> None:
    path = _build_docx(tmp_path / "queued.docx")
    idx = _make_indexer(tmp_path, skip_ocr=False)
    idx.qdrant_timeout_sec = 5

    idx.process_file(path)

    payloads = _content_payloads(idx)
    assert any(SAMPLE_TEXT in p["text"] and p["section"].startswith("image:") for p in payloads)
    assert idx.state_db.get_entry(str(path))["status"] == "ok"


@pytest.mark.parametrize("staged", [False, True])
def test_embedded_ocr_failure_is_not_a_successful_document(tmp_path, fake_ocr, monkeypatch, staged):
    import pytesseract
    path = _build_docx(tmp_path / "failed.docx")
    idx = _make_indexer(tmp_path, skip_ocr=False)
    idx.qdrant_timeout_sec = 5
    def fail(*args, **kwargs):
        raise RuntimeError("OCR unavailable")
    monkeypatch.setattr(pytesseract, "image_to_string", fail)
    if staged:
        assert idx.index_directory(stage="small")["error_files"] == 1
    else:
        with pytest.raises(RuntimeError, match="embedded_ocr_failed"):
            idx.process_file(path)
    assert idx.state_db.get_entry(str(path))["status"] == "error"
    assert idx.qdrant.points == []


def test_replacement_retries_same_checksum_after_failed_delete(tmp_path, fake_ocr):
    path = _build_docx(tmp_path / "retry.docx")
    idx = _make_indexer(tmp_path, skip_ocr=False)
    idx.qdrant_timeout_sec = 5
    idx.process_file(path)
    previous = list(idx.qdrant.points)
    calls = []
    def fail(*args, **kwargs):
        calls.append(True)
        raise RuntimeError("delete failed")
    idx._delete_file_vectors = fail
    for _ in range(2):
        with pytest.raises(RuntimeError, match="delete failed"):
            idx.process_file(path, fingerprint_override="new-checksum")
    assert len(calls) == 2
    assert idx.qdrant.points == previous
    assert idx.state_db.get_entry(str(path))["status"] == "error"


def test_unchanged_completed_file_does_not_delete_vectors(tmp_path, fake_ocr):
    path = _build_docx(tmp_path / "same.docx")
    idx = _make_indexer(tmp_path, skip_ocr=False)
    idx.qdrant_timeout_sec = 5
    idx.process_file(path)
    previous = list(idx.qdrant.points)
    def fail(*args, **kwargs):
        raise AssertionError("unchanged file must not be deleted")
    idx._delete_file_vectors = fail
    idx.process_file(path)
    assert idx.qdrant.points == previous
    assert all(p.get("modified") for p in _content_payloads(idx))


def test_indexer_disables_embedded_media_when_configured(tmp_path: Path, fake_ocr) -> None:
    path = _build_docx(tmp_path / "scan.docx")
    idx = _make_indexer(tmp_path, skip_ocr=False)
    idx.index_embedded_media = False

    idx.index_directory(stage="small")

    assert fake_ocr == []
    assert not any(str(p.get("section") or "").startswith("image:") for p in _content_payloads(idx))
    assert idx.state_db.get_entry(str(path))["status"] == "ok"


@pytest.mark.parametrize(
    ("method", "target"),
    [
        ("_extract_docx_document", "extract_docx_document"),
        ("_extract_pptx_document", "extract_pptx_document"),
        ("_extract_spreadsheet_document", "extract_spreadsheet_document"),
    ],
)
def test_office_extractors_receive_ocr_settings_from_indexer(monkeypatch, tmp_path: Path, method, target) -> None:
    seen: dict = {}

    def _fake(filepath, **kwargs):
        seen.update(kwargs)
        return ExtractedDocument(blocks=())

    monkeypatch.setattr(index_rag_module, target, _fake)
    idx = _make_indexer(tmp_path, skip_ocr=True)
    idx.ocr_tesseract_cmd = r"X:\tools\tesseract.exe"
    idx.index_embedded_media = False

    getattr(idx, method)(tmp_path / "any.bin")

    assert seen["skip_ocr"] is True
    assert seen["tesseract_cmd"] == r"X:\tools\tesseract.exe"
    assert seen["include_embedded"] is False
    assert seen["max_chars"] == 0


def test_indexer_constructor_config_defaults() -> None:
    from rag_catalog.core.rag_core import DEFAULT_CONFIG

    assert DEFAULT_CONFIG["index_embedded_media"] is True
    assert DEFAULT_CONFIG["index_cleanup_skip_failed_ratio"] == 0.1
    assert DEFAULT_CONFIG["indexer_heartbeat_path"] == "data/indexer_heartbeat.json"


# ── ocr_deferral helpers ──────────────────────────────────────────────────────


def test_deferred_ocr_candidate_rules() -> None:
    assert is_deferred_ocr_candidate(".pdf", "deferred_ocr")
    assert is_deferred_ocr_candidate(".pdf", "empty")
    assert is_deferred_ocr_candidate(".JPG", "empty")
    assert is_deferred_ocr_candidate(".docx", "deferred_ocr")
    assert not is_deferred_ocr_candidate(".docx", "empty")  # пустой docx без картинок — обычный backoff
    assert not is_deferred_ocr_candidate(".txt", "deferred_ocr")


def test_document_has_deferred_embedded_ocr() -> None:
    assert document_has_deferred_embedded_ocr(
        ExtractedDocument(blocks=(), metadata={"embedded_images": 2, "ocr_skipped": True})
    )
    assert not document_has_deferred_embedded_ocr(
        ExtractedDocument(blocks=(), metadata={"embedded_images": 0, "ocr_skipped": True})
    )
    assert not document_has_deferred_embedded_ocr(ExtractedDocument(blocks=(), metadata={"embedded_images": 2}))
    assert not document_has_deferred_embedded_ocr(None)


# ── п.2: block.section / block_kind в payload ────────────────────────────────


def test_chunk_provenance_prefers_explicit_block_section(tmp_path: Path) -> None:
    idx = _make_indexer(tmp_path)
    block = TextBlock(
        text="1. ЗАГОЛОВОК ИЗ ТЕКСТА\nтекст скана",
        section="image:word/media/image1.png",
        metadata={"source": "embedded_image", "member": "word/media/image1.png"},
    )

    prov = idx._chunk_provenance(chunk=block.text, chunk_index=0, doc_id="doc", block=block)

    assert prov["section"] == "image:word/media/image1.png"
    assert prov["block_kind"] == "embedded_image"
    assert prov["provenance"]["section"] == "image:word/media/image1.png"
    assert prov["provenance"]["block_kind"] == "embedded_image"


@pytest.mark.parametrize(
    ("section", "metadata", "kind"),
    [
        ("embedded:word/embeddings/x.xlsx", {"source": "embedded_document"}, "embedded_document"),
        ("embedded:word/embeddings/x.xlsx", {}, "embedded_document"),
        ("header", {}, "header"),
        ("footer", {}, "footer"),
        ("footnotes", {}, "footnotes"),
        ("textbox", {}, "textbox"),
        ("image:xl/media/image2.jpg", {}, "embedded_image"),
    ],
)
def test_chunk_provenance_block_kinds(tmp_path: Path, section, metadata, kind) -> None:
    idx = _make_indexer(tmp_path)
    block = TextBlock(text="Содержимое блока", section=section, metadata=metadata)

    prov = idx._chunk_provenance(chunk=block.text, chunk_index=3, doc_id="doc", block=block)

    assert prov["section"] == section
    assert prov["block_kind"] == kind


def test_chunk_provenance_falls_back_to_title_without_explicit_section(tmp_path: Path) -> None:
    idx = _make_indexer(tmp_path)
    chunk = "2.1 Технические характеристики\nМасса 3400 кг"

    without_block = idx._chunk_provenance(chunk=chunk, chunk_index=0, doc_id="doc", block=None)
    with_plain_block = idx._chunk_provenance(
        chunk=chunk, chunk_index=0, doc_id="doc", block=TextBlock(text=chunk, page=2)
    )

    assert without_block["section"] == "2.1 Технические характеристики"
    assert without_block["block_kind"] == "body"
    assert with_plain_block["section"] == "2.1 Технические характеристики"
    assert with_plain_block["block_kind"] == "body"
    assert with_plain_block["page"] == 2


def test_structured_chunking_keeps_section_of_merged_block(tmp_path: Path) -> None:
    idx = _make_indexer(tmp_path)
    document = ExtractedDocument(
        blocks=(
            TextBlock(text="Тело документа " * 5),
            TextBlock(text="Верхний колонтитул организации", section="header"),
            TextBlock(text="Текст со скана " * 5, section="image:word/media/image1.png",
                      metadata={"source": "embedded_image"}),
        )
    )

    items = idx._chunk_text_with_provenance(document)
    sections = [item["block"].section for item in items]

    assert "" in sections
    assert "header" in sections
    assert "image:word/media/image1.png" in sections
    image_item = next(item for item in items if item["block"].section.startswith("image:"))
    assert idx._block_kind(image_item["block"]) == "embedded_image"
